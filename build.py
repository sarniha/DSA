#!/usr/bin/env python3
"""Generate SIH PPT: AI-powered Predictive Analytics & Early Warning System
for the PAIMANA / MoSPI infrastructure project-monitoring platform.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette
NAVY      = RGBColor(0x0F, 0x2B, 0x46)
NAVY_DK   = RGBColor(0x08, 0x1B, 0x2E)
TEAL      = RGBColor(0x14, 0x74, 0x9C)
BLUE      = RGBColor(0x2E, 0x86, 0xAB)
SKY       = RGBColor(0x41, 0xA6, 0xC9)
ORANGE    = RGBColor(0xE8, 0x8A, 0x1A)
AMBER     = RGBColor(0xF2, 0xA9, 0x3B)
RED       = RGBColor(0xC0, 0x39, 0x2B)
GREEN     = RGBColor(0x2E, 0x8B, 0x57)
LIGHT     = RGBColor(0xF2, 0xF6, 0xFA)
LIGHT2    = RGBColor(0xE6, 0xEE, 0xF5)
TEXT      = RGBColor(0x1F, 0x2F, 0x3E)
GRAY      = RGBColor(0x6B, 0x7A, 0x8D)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

PROJECT = "InfraPredict AI"
SUB = "Predictive Analytics & Early Warning System for Infrastructure Project Monitoring"

# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(blank)

def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp

def tb(s, x, y, w, h, text, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.0, space_after=4, wrap=True):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return box

def shape_text(s, shp, text, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, font=FONT):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return shp

def bg(s):
    rect(s, 0, 0, SW, SH, fill=WHITE)

def header(s, title, tag=None, accent=TEAL):
    rect(s, 0, 0, SW, Inches(0.10), fill=NAVY)
    tb(s, Inches(0.55), Inches(0.30), Inches(10.8), Inches(0.7), title,
       size=30, color=NAVY, bold=True)
    if tag:
        tb(s, Inches(0.55), Inches(0.95), Inches(11.5), Inches(0.4), tag,
           size=14, color=accent, bold=False)
    rect(s, Inches(0.55), Inches(0.92) if tag else Inches(0.80), Inches(2.0), Inches(0.05), fill=accent)

def footer(s, n, note=PROJECT):
    rect(s, 0, Inches(7.14), SW, Inches(0.36), fill=LIGHT)
    tb(s, Inches(0.55), Inches(7.18), Inches(8), Inches(0.3), note, size=11, color=GRAY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(12.2), Inches(7.18), Inches(0.6), Inches(0.3), str(n), size=11, color=GRAY, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def card(s, x, y, w, h, title, body, accent=TEAL, tsize=15, bsize=12, tcolor=NAVY):
    rect(s, x, y, w, h, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    rect(s, x, y, Inches(0.09), h, fill=accent)
    tb(s, x+Inches(0.28), y+Inches(0.16), w-Inches(0.5), Inches(0.5), title, size=tsize, color=tcolor, bold=True)
    tb(s, x+Inches(0.28), y+Inches(0.62), w-Inches(0.5), h-Inches(0.7), body, size=bsize, color=TEXT, bold=False, line_spacing=1.1, space_after=3)

def chip(s, x, y, w, h, text, fill=TEAL, tsize=12, tcolor=WHITE):
    shp = rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.4)
    return shape_text(s, shp, text, size=tsize, color=tcolor, bold=True, align=PP_ALIGN.CENTER)

def bullet(s, x, y, w, items, size=14, gap=14, color=TEXT, bullet_char="\u25B8", bcolor=TEAL, h=Inches(3.0)):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15; p.space_after = Pt(gap)
        r = p.add_run(); r.text = bullet_char + "  "
        r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = bcolor; r.font.name = FONT
        r2 = p.add_run(); r2.text = it
        r2.font.size = Pt(size); r2.font.bold = False; r2.font.color.rgb = color; r2.font.name = FONT
    return box

# ================================================================ SLIDE 1: TITLE
s = slide(); bg(s)
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, 0, Inches(4.7), SH, fill=NAVY_DK)
# decorative
for i, (c, y) in enumerate([(TEAL, 5.6), (SKY, 6.1), (ORANGE, 6.6)]):
    rect(s, 0, Inches(y), Inches(4.7), Inches(0.08), fill=c)
# left brand
rect(s, Inches(0.6), Inches(0.7), Inches(0.09), Inches(1.1), fill=ORANGE)
tb(s, Inches(0.85), Inches(0.68), Inches(3.6), Inches(1.1),
   "Smart India Hackathon\n2026", size=20, color=WHITE, bold=True, line_spacing=1.05)
rect(s, Inches(0.6), Inches(1.95), Inches(3.4), Inches(0.02), fill=RGBColor(0x30,0x4a,0x66))
tb(s, Inches(0.6), Inches(2.15), Inches(3.9), Inches(1.2),
   "Problem Statement ID: 26103", size=18, color=SKY, bold=True)
tb(s, Inches(0.6), Inches(2.65), Inches(3.9), Inches(1.2),
   "Use case on web-based integrated\nproject-monitoring platform", size=13, color=RGBColor(0xB8,0xC7,0xD6), line_spacing=1.2)
tb(s, Inches(0.6), Inches(4.05), Inches(3.9), Inches(1.2),
   "AI for Infrastructure Monitoring", size=13, color=AMBER, bold=True, line_spacing=1.2)
tb(s, Inches(0.6), Inches(4.55), Inches(3.9), Inches(1.2),
   "MoSPI \u00b7 IPMD \u00b7 PAIMANA", size=12, color=RGBColor(0x8A,0xA0,0xB5), line_spacing=1.2)

# right title
rect(s, Inches(5.3), Inches(1.9), Inches(0.16), Inches(1.6), fill=ORANGE)
tb(s, Inches(5.6), Inches(1.85), Inches(7.3), Inches(1.7),
   "InfraPredict AI", size=48, color=WHITE, bold=True)
tb(s, Inches(5.6), Inches(2.55), Inches(7.3), Inches(1.3),
   "AI-Powered Predictive Analytics & Early Warning System\nfor Central Sector Infrastructure Project Monitoring",
   size=18, color=RGBColor(0xC9,0xD8,0xE6), line_spacing=1.3)
tb(s, Inches(5.6), Inches(3.9), Inches(7.3), Inches(1.2),
   "From Descriptive Monitoring \u2192 Predictive & Prescriptive Decision Support",
   size=15, color=SKY, bold=True)
tb(s, Inches(5.6), Inches(4.5), Inches(7.3), Inches(1.4),
   "Predict cost overruns \u00b7 Forecast time overruns \u00b7 Score project risk\n\u00b7 Issue early warnings \u00b7 Power an LLM-enabled Project Intelligence Assistant",
   size=13, color=RGBColor(0xA7,0xB8,0xC9), line_spacing=1.4)

rect(s, Inches(5.6), Inches(6.15), Inches(6.8), Inches(0.02), fill=RGBColor(0x2f,0x49,0x65))
tb(s, Inches(5.6), Inches(6.35), Inches(7.3), Inches(0.7),
   "Team Aurora   |   Open-Source AI/ML Stack   |   2026", size=13, color=RGBColor(0x8A,0xA0,0xB5), bold=True)

# ================================================================ SLIDE 2: BACKGROUND
s = slide(); bg(s)
header(s, "Background & Context", "The data problem: a 20-year national infrastructure database", TEAL)
# 3 stat cards
stats = [
    ("1,981", "Ongoing infrastructure projects", "across 17 Central Ministries/Departments & 22 sectors"),
    ("\u20b937.13 L cr", "Aggregate original cost", "revised cost \u20b942.78 L cr"),
    ("\u20b920.36 L cr", "Cumulative expenditure", "monitored on a monthly cycle"),
]
x = Inches(0.55)
for val, lab, sub in stats:
    shp = rect(s, x, Inches(1.55), Inches(4.05), Inches(1.9), fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    rect(s, x, Inches(1.55), Inches(4.05), Inches(0.12), fill=TEAL)
    tb(s, x, Inches(1.8), Inches(4.05), Inches(0.6), val, size=34, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    tb(s, x, Inches(2.5), Inches(4.05), Inches(0.5), lab, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    tb(s, x, Inches(2.95), Inches(4.05), Inches(0.5), sub, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    x += Inches(4.25)

card(s, Inches(0.55), Inches(3.8), Inches(6.1), Inches(3.1), "Who monitors & how",
     "The Infrastructure & Project Monitoring Division (IPMD), Ministry of Statistics & Programme Implementation (MoSPI), monitors Central Sector Infrastructure Projects costing \u20b9150 crore and above across all infrastructural Ministries/Departments.\n\nMonitored through the Online Computerised Monitoring System (OCMS) since 2006, then modernised to the PAIMANA portal \u2014 Project Assessment, Infrastructure Monitoring and Analytics for Nation-building.",
     accent=TEAL, tsize=16, bsize=12)

card(s, Inches(6.85), Inches(3.8), Inches(6.0), Inches(3.1), "The data ecosystem",
     "PAIMANA is a web-based integrated project-monitoring platform acting as a national repository of infrastructure projects.\n\nIt captures approved cost, revised cost, expenditure, implementation timelines, physical progress, milestones, implementing agencies and project status, updated monthly through role-based access and APIs.",
     accent=ORANGE, tsize=16, bsize=12)
footer(s, 2)

# ================================================================ SLIDE 3: THE CHALLENGE
s = slide(); bg(s)
header(s, "The Problem", "Despite rich data, overruns persist \u2014 monitoring is reactive, not predictive", ORANGE)
bullet(s, Inches(0.55), Inches(1.7), Inches(6.0), [
 "Cost overruns: approved cost revises from \u20b937.13 to \u20b942.78 lakh crore (+15%).",
 "Time overruns: delays in milestone achievement push assets beyond scheduled completion.",
 "Contractual & implementation bottlenecks, resource constraints and execution risks.",
 "Completion of public assets is delayed and costs escalate significantly.",
], size=14, gap=16, bcolor=ORANGE)

card(s, Inches(6.9), Inches(1.7), Inches(5.95), Inches(3.3), "What PAIMANA gives us today",
     "\u2713 Strong descriptive & reporting capabilities\n\u2713 Role-based access, monthly updates, APIs\n\n\u2717 No forecasting of future overruns\n\u2717 No project-level risk score or early warning\n\u2717 No automated driver/root-cause analysis\n\u2717 No natural-language querying of project data",
     accent=TEAL, tsize=15, bsize=13)

rect(s, Inches(0.55), Inches(4.3), Inches(12.3), Inches(1.05), fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
tb(s, Inches(0.9), Inches(4.55), Inches(11.7), Inches(0.7),
   "The gap: move from descriptive monitoring to a PREDICTIVE & PRESCRIPTIVE decision-support system.",
   size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

rect(s, Inches(0.55), Inches(5.6), Inches(12.3), Inches(1.35), fill=LIGHT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
tb(s, Inches(0.85), Inches(5.75), Inches(11.8), Inches(0.5), "Why now?", size=15, color=NAVY, bold=True)
tb(s, Inches(0.85), Inches(6.15), Inches(11.8), Inches(0.7),
   "Nearly two decades of OCMS history + continuously updated PAIMANA data = a large, diverse, real-time dataset ideal for AI/ML and LLMs to forecast cost escalation, schedule delays and implementation risk before they materialise.",
   size=13, color=TEXT, line_spacing=1.15)
footer(s, 3)

# ================================================================ SLIDE 4: PROBLEM STATEMENT
s = slide(); bg(s)
header(s, "Problem Statement & Scope of Work", "The three technical dimensions we must address", BLUE)
card(s, Inches(0.55), Inches(1.7), Inches(4.0), Inches(4.6), "A) Predictive Modelling",
     "Develop & evaluate statistical and AI/ML models (open-source) to forecast cost overruns, time overruns and implementation risk using historical + live project data.",
     accent=TEAL, tsize=16, bsize=13)
card(s, Inches(4.75), Inches(1.7), Inches(4.0), Inches(4.6), "B) AI vs Conventional",
     "Assess whether AI/ML delivers significant gains over conventional statistical methods in prediction accuracy, early-warning capability and decision support.",
     accent=ORANGE, tsize=16, bsize=13)
card(s, Inches(8.95), Inches(1.7), Inches(3.9), Inches(4.6), "C) Feature Attribution",
     "Build models from existing CUF fields and measure how much predictive power comes from current CUF variables versus additional variables not yet captured.",
     accent=GREEN, tsize=16, bsize=13)
for i, txt in enumerate(["A", "B", "C"]):
    cx = [Inches(2.55), Inches(6.75), Inches(10.95)][i]
    shape_text(s, rect(s, cx, Inches(1.55), Inches(0.55), Inches(0.55), fill=NAVY, shape=MSO_SHAPE.OVAL), txt, size=20)
rect(s, Inches(0.55), Inches(6.4), Inches(12.3), Inches(0.6), fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
tb(s, Inches(0.9), Inches(6.45), Inches(11.8), Inches(0.5),
   "Deliverable: a deployed, open-source AI-powered monitoring dashboard + risk scoring + early-warning + LLM assistant.",
   size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 4)

# ================================================================ SLIDE 5: SOLUTION OVERVIEW
s = slide(); bg(s)
header(s, "Proposed Solution", "InfraPredict AI \u2014 an end-to-end predictive & prescriptive monitoring suite", TEAL)
tb(s, Inches(0.55), Inches(1.7), Inches(12.3), Inches(0.5),
   "Five integrated modules that transform raw PAIMANA data into actionable, evidence-based decisions.",
   size=15, color=GRAY, align=PP_ALIGN.CENTER)

mods = [
    ("Predictive Models", "Forecast cost & time overrun probability and magnitude per project.", TEAL),
    ("Risk Scoring", "Composite 0-100 project risk score with sector & agency context.", ORANGE),
    ("Early-Warning Alerts", "Pre-defined thresholds trigger alerts/notifications before overruns materialise.", BLUE),
    ("Driver Analysis", "Shapley-style attribution of the factors driving each project's escalation.", GREEN),
    ("AI Dashboard + LLM Assistant", "Interactive analytics & natural-language querying of the portfolio.", NAVY),
]
x = Inches(0.55); w = Inches(2.42)
for t, d, c in mods:
    shp = rect(s, x, Inches(2.4), w, Inches(2.5), fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    rect(s, x, Inches(2.4), w, Inches(0.13), fill=c)
    shape_text(s, rect(s, x + w/2 - Inches(0.45), Inches(2.7), Inches(0.9), Inches(0.9), fill=c, shape=MSO_SHAPE.OVAL), t.split()[0][0], size=26, bold=True)
    tb(s, x+Inches(0.2), Inches(3.72), w-Inches(0.4), Inches(0.6), t, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    tb(s, x+Inches(0.25), Inches(4.25), w-Inches(0.5), Inches(0.6), d, size=11, color=TEXT, align=PP_ALIGN.CENTER, line_spacing=1.12)
    x += Inches(2.5)

rect(s, Inches(0.55), Inches(5.2), Inches(12.3), Inches(1.75), fill=LIGHT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
tb(s, Inches(0.85), Inches(5.35), Inches(11.8), Inches(0.5), "How it flows", size=15, color=NAVY, bold=True)
flow = [("Collect", TEAL), ("Engineer", BLUE), ("Predict", ORANGE), ("Score & Alert", RED), ("Act & Monitor", GREEN)]
fx = Inches(1.0)
for t, c in flow:
    shp = rect(s, fx, Inches(5.9), Inches(1.9), Inches(0.6), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    shape_text(s, shp, t, size=14)
    fx += Inches(2.15)
    if t != "Act & Monitor":
        shape_text(s, rect(s, fx-Inches(0.32), Inches(6.0), Inches(0.3), Inches(0.4), fill=GRAY, shape=MSO_SHAPE.RIGHT_ARROW), "", size=1)
footer(s, 5)

# ---------------------------------------------------------------- save
import os
OUT_PATH = "/mnt/user-data/outputs/InfraPredict_AI.pptx"
os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
prs.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")